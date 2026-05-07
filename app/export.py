from __future__ import annotations

import io
import json
from datetime import datetime, timezone
from typing import Union

import pandas as pd

from app.models import (
    CapitalPanel,
    DependencyReport,
    EnrichedReport,
    MaterialityScore,
)

_TOOL_VERSION = "0.1.0"

AnyReport = Union[DependencyReport, EnrichedReport]

# RAG colour fills for Excel (openpyxl ARGB hex, no '#')
_FILL_HIGH = "FFD32F2F"    # red
_FILL_MEDIUM = "FFF57C00"  # amber
_FILL_LOW = "FF388E3C"     # green
_FILL_NA = "FF9E9E9E"      # grey


def _get_base_report(report: AnyReport) -> DependencyReport:
    if isinstance(report, EnrichedReport):
        return report.base_report
    return report


def flatten_report(report: AnyReport) -> pd.DataFrame:
    """Flatten a DependencyReport/EnrichedReport into a single DataFrame."""
    base = _get_base_report(report)
    rows = []

    for panel in [base.natural_capital, base.social_capital, base.human_capital]:
        for scored in panel.dependencies:
            dep = scored.dependency
            rows.append(
                {
                    "capital_type": panel.capital_type.value,
                    "category": dep.category,
                    "subcategory": dep.subcategory,
                    "dependency_id": dep.id,
                    "label": dep.label,
                    "materiality_score": scored.materiality_score.value,
                    "materiality_numeric": scored.materiality_numeric,
                    "sasb_basis": scored.sasb_basis,
                    "rationale": scored.rationale,
                    "llm_adjusted": scored.llm_adjusted,
                    "llm_adjustment_note": scored.llm_adjustment_note or "",
                    "tnfd_ecosystem_service": dep.tnfd_ecosystem_service or "",
                    "sasb_disclosure_topic": dep.sasb_disclosure_topic,
                    "indicators": "; ".join(dep.indicators),
                }
            )

    # Append LLM freetext additions if present
    if isinstance(report, EnrichedReport):
        for ft in report.llm_additions:
            rows.append(
                {
                    "capital_type": ft.capital_type.value,
                    "category": "LLM Identified",
                    "subcategory": "LLM Identified",
                    "dependency_id": "LLM",
                    "label": ft.label,
                    "materiality_score": ft.materiality_score.value,
                    "materiality_numeric": {
                        MaterialityScore.HIGH: 3,
                        MaterialityScore.MEDIUM: 2,
                        MaterialityScore.LOW: 1,
                    }.get(ft.materiality_score, 1),
                    "sasb_basis": "LLM analysis",
                    "rationale": ft.llm_rationale,
                    "llm_adjusted": True,
                    "llm_adjustment_note": "LLM-identified dependency",
                    "tnfd_ecosystem_service": "",
                    "sasb_disclosure_topic": "",
                    "indicators": "",
                }
            )

    return pd.DataFrame(rows)


def _build_metadata(report: AnyReport) -> dict:
    base = _get_base_report(report)
    llm_enriched = isinstance(report, EnrichedReport) and report.llm_enriched
    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "tool_version": _TOOL_VERSION,
        "frameworks_used": base.frameworks_used,
        "llm_enriched": llm_enriched,
        "llm_model": report.llm_model_used if isinstance(report, EnrichedReport) else None,
        "sasb_industry_code": base.sasb_mapping.sasb_industry_code,
        "sasb_confidence": base.sasb_mapping.confidence,
    }


def to_json(report: AnyReport) -> str:
    base = _get_base_report(report)
    llm_enriched = isinstance(report, EnrichedReport) and report.llm_enriched

    payload = {
        "metadata": _build_metadata(report),
        "company": base.company.model_dump(exclude={"raw_yfinance_data"}),
        "sasb_mapping": base.sasb_mapping.model_dump(),
        "dependencies": {
            "natural_capital": [
                _scored_dep_to_dict(s) for s in base.natural_capital.dependencies
            ],
            "social_capital": [
                _scored_dep_to_dict(s) for s in base.social_capital.dependencies
            ],
            "human_capital": [
                _scored_dep_to_dict(s) for s in base.human_capital.dependencies
            ],
        },
    }

    if isinstance(report, EnrichedReport) and report.llm_enriched:
        payload["llm_enrichment"] = {
            "narrative": report.llm_narrative,
            "evidence_sources": report.llm_evidence_sources,
            "llm_identified_dependencies": [
                ft.model_dump() for ft in report.llm_additions
            ],
        }

    return json.dumps(payload, indent=2, default=str)


def _scored_dep_to_dict(scored) -> dict:
    return {
        "id": scored.dependency.id,
        "label": scored.dependency.label,
        "category": scored.dependency.category,
        "subcategory": scored.dependency.subcategory,
        "materiality_score": scored.materiality_score.value,
        "materiality_numeric": scored.materiality_numeric,
        "rationale": scored.rationale,
        "llm_adjusted": scored.llm_adjusted,
        "llm_adjustment_note": scored.llm_adjustment_note,
        "sasb_disclosure_topic": scored.dependency.sasb_disclosure_topic,
        "tnfd_ecosystem_service": scored.dependency.tnfd_ecosystem_service,
        "indicators": scored.dependency.indicators,
    }


def to_csv(report: AnyReport) -> bytes:
    df = flatten_report(report)
    return df.to_csv(index=False).encode("utf-8")


def to_excel(report: AnyReport) -> bytes:
    """Export to Excel with conditional formatting on the materiality column."""
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter

    df = flatten_report(report)

    wb = Workbook()
    ws = wb.active
    ws.title = "Capital Dependencies"

    # Header row
    header_font = Font(bold=True)
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Data rows with conditional fill on materiality_score column
    mat_col_idx = list(df.columns).index("materiality_score") + 1  # 1-based

    fill_map = {
        "high": PatternFill("solid", fgColor=_FILL_HIGH),
        "medium": PatternFill("solid", fgColor=_FILL_MEDIUM),
        "low": PatternFill("solid", fgColor=_FILL_LOW),
    }

    for row_data in df.itertuples(index=False):
        row_vals = list(row_data)
        ws.append(row_vals)
        row_idx = ws.max_row
        score_cell = ws.cell(row=row_idx, column=mat_col_idx)
        fill = fill_map.get(str(score_cell.value).lower())
        if fill:
            score_cell.fill = fill

    # Auto-width for key columns
    for col_idx, col_name in enumerate(df.columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(col_name)), df[col_name].astype(str).str.len().max())
        ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
