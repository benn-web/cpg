"""Tests for export.py."""
from __future__ import annotations

import json

import pandas as pd
import pytest

from app.export import flatten_report, to_csv, to_excel, to_json
from app.models import DependencyReport


def test_to_json_valid_structure(sample_report):
    output = to_json(sample_report)
    data = json.loads(output)
    assert "metadata" in data
    assert "company" in data
    assert "dependencies" in data
    assert "sasb_mapping" in data


def test_to_json_metadata_fields(sample_report):
    data = json.loads(to_json(sample_report))
    meta = data["metadata"]
    assert "generated_at" in meta
    assert "frameworks_used" in meta
    assert "tool_version" in meta
    assert "llm_enriched" in meta
    frameworks = meta["frameworks_used"]
    assert any("TNFD" in f for f in frameworks)
    assert any("SASB" in f for f in frameworks)


def test_to_json_has_all_capital_types(sample_report):
    data = json.loads(to_json(sample_report))
    deps = data["dependencies"]
    assert "natural_capital" in deps
    assert "social_capital" in deps
    assert "human_capital" in deps


def test_to_json_excludes_raw_yfinance_data(sample_report):
    output = to_json(sample_report)
    assert "raw_yfinance_data" not in output


def test_to_csv_correct_columns(sample_report):
    csv_bytes = to_csv(sample_report)
    df = pd.read_csv(pd.io.common.BytesIO(csv_bytes))
    expected_cols = {
        "capital_type", "category", "label",
        "materiality_score", "materiality_numeric", "rationale",
    }
    assert expected_cols.issubset(set(df.columns))


def test_to_csv_row_count(sample_report):
    csv_bytes = to_csv(sample_report)
    df = pd.read_csv(pd.io.common.BytesIO(csv_bytes))
    total_deps = (
        len(sample_report.natural_capital.dependencies)
        + len(sample_report.social_capital.dependencies)
        + len(sample_report.human_capital.dependencies)
    )
    assert len(df) == total_deps


def test_to_excel_valid_workbook(sample_report):
    import io
    from openpyxl import load_workbook
    excel_bytes = to_excel(sample_report)
    wb = load_workbook(io.BytesIO(excel_bytes))
    assert "Capital Dependencies" in wb.sheetnames


def test_flatten_report_numeric_scores(sample_report):
    df = flatten_report(sample_report)
    high_rows = df[df["materiality_score"] == "high"]
    assert all(high_rows["materiality_numeric"] == 3)
    medium_rows = df[df["materiality_score"] == "medium"]
    assert all(medium_rows["materiality_numeric"] == 2)
    low_rows = df[df["materiality_score"] == "low"]
    assert all(low_rows["materiality_numeric"] == 1)


def test_flatten_report_all_capital_types(sample_report):
    df = flatten_report(sample_report)
    types = set(df["capital_type"].unique())
    assert "natural" in types
    assert "social" in types
    assert "human" in types
