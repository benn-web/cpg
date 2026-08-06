#!/usr/bin/env python3
"""Build the consultancy income forecast workbook (Income_Forecast.xlsx).

Percentage-of-completion revenue recognition for confirmed engagements.
Run:  python tools/build_income_forecast.py   (writes ./Income_Forecast.xlsx)

Note: openpyxl writes formulas without cached results, so the freshly generated
file shows blank formula cells until opened in Excel/LibreOffice, which recalculate
on open. The committed Income_Forecast.xlsx additionally has verified cached values.
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

# ---- palette / fonts ----
ARIAL = "Arial"
F_TITLE   = Font(name=ARIAL, size=16, bold=True, color="1F3864")
F_SUB     = Font(name=ARIAL, size=10, italic=True, color="595959")
F_HDR     = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_BODY    = Font(name=ARIAL, size=10, color="000000")          # formula / black
F_INPUT   = Font(name=ARIAL, size=10, color="0000FF")          # user input / blue
F_BOLD    = Font(name=ARIAL, size=10, bold=True, color="000000")
F_LINK    = Font(name=ARIAL, size=10, color="008000")          # cross-sheet / green

FILL_HDR   = PatternFill("solid", fgColor="1F3864")
FILL_HDR2  = PatternFill("solid", fgColor="2E5496")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")            # pale yellow = fill me in
FILL_CALC  = PatternFill("solid", fgColor="F2F2F2")
FILL_TOTAL = PatternFill("solid", fgColor="D9E1F2")
FILL_RED   = PatternFill("solid", fgColor="F8CBAD")
FILL_AMBER = PatternFill("solid", fgColor="FFE699")
FILL_GREEN = PatternFill("solid", fgColor="C6E0B4")
FONT_RED   = Font(name=ARIAL, size=10, color="843C0C")
FONT_AMBER = Font(name=ARIAL, size=10, color="806000")

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT   = Alignment(horizontal="left", vertical="center")
WRAP   = Alignment(horizontal="left", vertical="top", wrap_text=True)

GBP   = '£#,##0;(£#,##0);-'
GBP0  = '£#,##0'
PCT   = '0.0%'
DATEF = 'MMM-YYYY'

wb = Workbook()

def hdr(ws, row, headers, start=1, fill=FILL_HDR):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start+i, value=h)
        c.font = F_HDR; c.fill = fill; c.alignment = CENTER; c.border = BORDER

# ============================================================ INSTRUCTIONS
ws = wb.active
ws.title = "Instructions"
ws.sheet_view.showGridLines = False
ws["A1"] = "Consultancy Income Forecast"; ws["A1"].font = F_TITLE
ws["A2"] = "Percentage-of-completion revenue recognition for confirmed engagements · GBP (£) · Financial year September–August"
ws["A2"].font = F_SUB
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 110

lines = [
    ("How this workbook works", "h"),
    ("Every engagement is confirmed work with a Matter Code and an agreed fixed fee. Each month the engagement "
     "manager (EM) sets the cumulative % complete; the workbook recognises the corresponding slice of the fee. "
     "Actual net fee income (NFI) is pulled automatically from the Month-End Provisions tab and compared against "
     "the forecast, and any variance is flagged.", "p"),
    ("The five tabs", "h"),
    ("1.  Matters Register — master list of engagements (Matter Code, client, EM, total fixed fee, dates, status). "
     "Add a row here first for every new confirmed engagement.", "p"),
    ("2.  Forecast Log — the working record. One row per engagement per month. You enter the Matter Code, the Month, "
     "and the cumulative % complete. Everything else calculates. This is the permanent, append-only history — never "
     "overwrite a past month's row; add a new row for the new month.", "p"),
    ("3.  Month-End Provisions — paste your monthly NFI provisions extract here (Matter Code, Month, NFI Amount). "
     "The Forecast Log reads actuals from this tab by matching Matter Code + Month.", "p"),
    ("4.  Dashboard — portfolio roll-up by month, by engagement and by EM, with a forecast-vs-actual chart.", "p"),
    ("5.  Instructions — this page.", "p"),
    ("Monthly routine", "h"),
    ("a.  Paste the latest NFI provisions into Month-End Provisions.", "p"),
    ("b.  In the Forecast Log, add one new row per active engagement for the new month and enter the cumulative % complete.", "p"),
    ("c.  Review the Variance column. A red cell = actual came in UNDER forecast (under-variance); an amber cell = "
     "actual came in OVER forecast (over-variance).", "p"),
    ("d.  Re-allocate next month: use the Cumulative Actual and Remaining Fee columns to re-base your % complete so the "
     "remaining fee is spread across the remaining months. The workbook flags the variance; the EM decides how to re-spread it.", "p"),
    ("Colour key", "h"),
    ("Blue text on pale-yellow fill = cells you fill in.   Black text = calculated, do not edit.   "
     "Green text = pulled from another tab.   Red / amber fills = under / over variance flags.", "p"),
    ("Notes", "h"),
    ("· % complete is stored as a fraction (enter 25% or 0.25, not 25).   · All amounts are GBP.   "
     "· Dates are the first of the month (e.g. 01-Sep-2026 shows as Sep-2026).   "
     "· The three demo engagements (M-1001/1002/1003) are examples — delete them once you add your own.", "p"),
]
r = 4
for text, kind in lines:
    c = ws.cell(row=r, column=2, value=text)
    if kind == "h":
        c.font = Font(name=ARIAL, size=12, bold=True, color="1F3864"); r += 1
    else:
        c.font = F_BODY; c.alignment = WRAP
        ws.row_dimensions[r].height = 15 * (1 + len(text)//95)
        r += 1
    r += 0

# ============================================================ MATTERS REGISTER
ws = wb.create_sheet("Matters Register")
ws.sheet_view.showGridLines = False
ws["A1"] = "Matters Register"; ws["A1"].font = F_TITLE
ws["A2"] = "One row per confirmed engagement. Blue cells are inputs. Add new engagements below the demo rows."
ws["A2"].font = F_SUB
reg_headers = ["Matter Code", "Client", "Engagement Name", "Engagement Manager",
               "Total Fee (£)", "Start Date", "Expected End Date", "Status"]
HDR_ROW = 4
hdr(ws, HDR_ROW, reg_headers)
reg_rows = [
    ["M-1001", "Aurora Retail Group", "Supply Chain Diagnostic", "Priya Shah", 120000,
     dt.date(2026,9,1), dt.date(2027,2,28), "Active"],
    ["M-1002", "Northwind Energy", "Cost Transformation", "James Okoro", 90000,
     dt.date(2026,10,1), dt.date(2027,3,31), "Active"],
    ["M-1003", "Meridian Health", "Operating Model Review", "Priya Shah", 60000,
     dt.date(2026,9,1), dt.date(2026,12,31), "Active"],
]
REG_FIRST = HDR_ROW + 1
REG_CAP = 40                      # capacity rows for lookups
for i, row in enumerate(reg_rows):
    rr = REG_FIRST + i
    for j, val in enumerate(row):
        c = ws.cell(row=rr, column=1+j, value=val)
        c.font = F_INPUT; c.fill = FILL_INPUT; c.border = BORDER
        if j == 4: c.number_format = GBP0
        if j in (5, 6): c.number_format = 'DD-MMM-YYYY'; c.alignment = CENTER
        if j == 0: c.alignment = CENTER
# blank input rows to capacity (styled, empty)
for rr in range(REG_FIRST + len(reg_rows), REG_FIRST + 20):
    for j in range(len(reg_headers)):
        c = ws.cell(row=rr, column=1+j)
        c.font = F_INPUT; c.fill = FILL_INPUT; c.border = BORDER
        if j == 4: c.number_format = GBP0
        if j in (5, 6): c.number_format = 'DD-MMM-YYYY'
widths = [12, 22, 26, 22, 14, 14, 16, 12]
for i, w in enumerate(widths):
    ws.column_dimensions[get_column_letter(1+i)].width = w
# status validation
dv_status = DataValidation(type="list", formula1='"Active,Complete,On hold"', allow_blank=True)
ws.add_data_validation(dv_status)
dv_status.add(f"H{REG_FIRST}:H{REG_FIRST+19}")
# fee non-negative
dv_fee = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
ws.add_data_validation(dv_fee)
dv_fee.add(f"E{REG_FIRST}:E{REG_FIRST+19}")
# Excel table (filter/sort + defined name + auto-expand on append)
_tbl = Table(displayName="tblMatters", ref=f"A{HDR_ROW}:H{REG_FIRST+19}")
_tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=False,
                                     showColumnStripes=False, showFirstColumn=False, showLastColumn=False)
ws.add_table(_tbl)

REG_MATTER = f"'Matters Register'!$A${REG_FIRST}:$A${REG_FIRST+REG_CAP-1}"
REG_FEE    = f"'Matters Register'!$E${REG_FIRST}:$E${REG_FIRST+REG_CAP-1}"
REG_EM     = f"'Matters Register'!$D${REG_FIRST}:$D${REG_FIRST+REG_CAP-1}"
REG_CLIENT = f"'Matters Register'!$B${REG_FIRST}:$B${REG_FIRST+REG_CAP-1}"

# ============================================================ MONTH-END PROVISIONS
ws = wb.create_sheet("Month-End Provisions")
ws.sheet_view.showGridLines = False
ws["A1"] = "Month-End Provisions (Net Fee Income)"; ws["A1"].font = F_TITLE
ws["A2"] = "Paste your monthly NFI provisions extract here. The Forecast Log pulls actuals from this tab by Matter Code + Month."
ws["A2"].font = F_SUB
prov_headers = ["Matter Code", "Month", "NFI Amount (£)"]
hdr(ws, HDR_ROW, prov_headers)
prov_rows = [
    ["M-1001", dt.date(2026,9,1), 12000],
    ["M-1001", dt.date(2026,10,1), 15000],
    ["M-1001", dt.date(2026,11,1), 26000],
    ["M-1002", dt.date(2026,10,1), 13500],
    ["M-1003", dt.date(2026,9,1), 12000],
    ["M-1003", dt.date(2026,10,1), 11000],
]
PROV_FIRST = HDR_ROW + 1
PROV_CAP = 1000
for i, row in enumerate(prov_rows):
    rr = PROV_FIRST + i
    for j, val in enumerate(row):
        c = ws.cell(row=rr, column=1+j, value=val)
        c.font = F_INPUT; c.fill = FILL_INPUT; c.border = BORDER
        if j == 0: c.alignment = CENTER
        if j == 1: c.number_format = DATEF; c.alignment = CENTER
        if j == 2: c.number_format = GBP0
for rr in range(PROV_FIRST + len(prov_rows), PROV_FIRST + 24):
    for j in range(len(prov_headers)):
        c = ws.cell(row=rr, column=1+j)
        c.font = F_INPUT; c.fill = FILL_INPUT; c.border = BORDER
        if j == 1: c.number_format = DATEF
        if j == 2: c.number_format = GBP0
for col, w in zip("ABC", [12, 14, 16]):
    ws.column_dimensions[col].width = w
# matter code validation (list from register)
dv_pm = DataValidation(type="list", formula1=f"={REG_MATTER}", allow_blank=True)
ws.add_data_validation(dv_pm)
dv_pm.add(f"A{PROV_FIRST}:A{PROV_FIRST+23}")
_tbl2 = Table(displayName="tblProvisions", ref=f"A{HDR_ROW}:C{PROV_FIRST+23}")
_tbl2.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=False,
                                      showColumnStripes=False, showFirstColumn=False, showLastColumn=False)
ws.add_table(_tbl2)

PLAST = PROV_FIRST + PROV_CAP - 1
PROV_MATTER = f"'Month-End Provisions'!$A${PROV_FIRST}:$A${PLAST}"
PROV_MONTH  = f"'Month-End Provisions'!$B${PROV_FIRST}:$B${PLAST}"
PROV_NFI    = f"'Month-End Provisions'!$C${PROV_FIRST}:$C${PLAST}"

# ============================================================ FORECAST LOG
ws = wb.create_sheet("Forecast Log")
ws.sheet_view.showGridLines = False
ws["A1"] = "Forecast Log — permanent monthly record"; ws["A1"].font = F_TITLE
ws["A2"] = ("Append one row per engagement per month. Enter Matter Code, Month and Cumulative % Complete (blue). "
            "Everything else calculates. Never overwrite a past month — add a new row.")
ws["A2"].font = F_SUB
log_headers = ["Matter Code", "Month", "Total Fee (£)", "Cum % Complete",
               "Forecast Cum Revenue (£)", "Forecast Recognised This Month (£)",
               "Actual NFI This Month (£)", "Variance (£)",
               "Cumulative Actual to Date (£)", "Remaining Fee (£)"]
hdr(ws, HDR_ROW, log_headers)
LOG_FIRST = HDR_ROW + 1
LOG_N = 60
LLAST = LOG_FIRST + LOG_N - 1
log_rows = [
    ["M-1001", dt.date(2026,9,1), 0.10],
    ["M-1001", dt.date(2026,10,1), 0.25],
    ["M-1001", dt.date(2026,11,1), 0.45],
    ["M-1002", dt.date(2026,10,1), 0.15],
    ["M-1002", dt.date(2026,11,1), 0.30],
    ["M-1003", dt.date(2026,9,1), 0.20],
    ["M-1003", dt.date(2026,10,1), 0.40],
    ["M-1003", dt.date(2026,11,1), 0.65],
]
E_col_range = f"$E${LOG_FIRST}:$E${LLAST}"
A_col_range = f"$A${LOG_FIRST}:$A${LLAST}"
B_col_range = f"$B${LOG_FIRST}:$B${LLAST}"

for i in range(LOG_N):
    rr = LOG_FIRST + i
    A = f"A{rr}"; B = f"B{rr}"; C = f"C{rr}"; D = f"D{rr}"; E = f"E{rr}"
    Fc = f"F{rr}"; G = f"G{rr}"; Hc = f"H{rr}"
    demo = log_rows[i] if i < len(log_rows) else None
    # inputs
    a = ws.cell(row=rr, column=1, value=(demo[0] if demo else None))
    a.font = F_INPUT; a.fill = FILL_INPUT; a.alignment = CENTER; a.border = BORDER
    b = ws.cell(row=rr, column=2, value=(demo[1] if demo else None))
    b.font = F_INPUT; b.fill = FILL_INPUT; b.number_format = DATEF; b.alignment = CENTER; b.border = BORDER
    d = ws.cell(row=rr, column=4, value=(demo[2] if demo else None))
    d.font = F_INPUT; d.fill = FILL_INPUT; d.number_format = PCT; d.alignment = CENTER; d.border = BORDER
    # C Total Fee (link to register)
    c = ws.cell(row=rr, column=3,
        value=f'=IF($A{rr}="","",IFERROR(INDEX({REG_FEE},MATCH($A{rr},{REG_MATTER},0)),"?"))')
    c.font = F_LINK; c.number_format = GBP0; c.border = BORDER
    # E Forecast Cum Revenue
    e = ws.cell(row=rr, column=5, value=f'=IF(OR($A{rr}="",$D{rr}=""),"",$D{rr}*$C{rr})')
    e.font = F_BODY; e.number_format = GBP; e.border = BORDER; e.fill = FILL_CALC
    # F Forecast recognised this month = cum revenue this month - cum revenue of prior month
    #   (prior month found with MAXIFS; references col E only, so no self-reference/circularity)
    prior = f'_xlfn.MAXIFS({B_col_range},{A_col_range},$A{rr},{B_col_range},"<"&$B{rr})'
    f = ws.cell(row=rr, column=6,
        value=(f'=IF($E{rr}="","",$E{rr}-SUMIFS({E_col_range},{A_col_range},$A{rr},'
               f'{B_col_range},"="&{prior}))'))
    f.font = F_BODY; f.number_format = GBP; f.border = BORDER; f.fill = FILL_CALC
    # G Actual NFI this month (pulled from provisions; blank if none posted)
    g = ws.cell(row=rr, column=7,
        value=(f'=IF($A{rr}="","",IF(COUNTIFS({PROV_MATTER},$A{rr},{PROV_MONTH},$B{rr})=0,"",'
               f'SUMIFS({PROV_NFI},{PROV_MATTER},$A{rr},{PROV_MONTH},$B{rr})))'))
    g.font = F_LINK; g.number_format = GBP; g.border = BORDER; g.fill = FILL_CALC
    # H Variance = actual - forecast (blank if no actual)
    h = ws.cell(row=rr, column=8, value=f'=IF(OR($G{rr}="",$F{rr}=""),"",$G{rr}-$F{rr})')
    h.font = F_BODY; h.number_format = GBP; h.border = BORDER
    # I Cumulative actual to date
    ii = ws.cell(row=rr, column=9,
        value=(f'=IF($A{rr}="","",SUMIFS({PROV_NFI},{PROV_MATTER},$A{rr},{PROV_MONTH},"<="&$B{rr}))'))
    ii.font = F_LINK; ii.number_format = GBP; ii.border = BORDER; ii.fill = FILL_CALC
    # J Remaining fee = total fee - cumulative actual
    jj = ws.cell(row=rr, column=10, value=f'=IF($A{rr}="","",$C{rr}-$I{rr})')
    jj.font = F_BODY; jj.number_format = GBP; jj.border = BORDER; jj.fill = FILL_CALC

log_widths = [12, 11, 13, 13, 16, 18, 16, 12, 16, 14]
for i, w in enumerate(log_widths):
    ws.column_dimensions[get_column_letter(1+i)].width = w
ws.row_dimensions[HDR_ROW].height = 42
for cc in range(1, len(log_headers)+1):
    ws.cell(row=HDR_ROW, column=cc).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
# validations
dv_lm = DataValidation(type="list", formula1=f"={REG_MATTER}", allow_blank=True)
ws.add_data_validation(dv_lm); dv_lm.add(f"A{LOG_FIRST}:A{LLAST}")
dv_pctl = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                         error="Enter cumulative % complete as a fraction between 0 and 1 (e.g. 0.25 for 25%).",
                         errorTitle="% complete")
ws.add_data_validation(dv_pctl); dv_pctl.add(f"D{LOG_FIRST}:D{LLAST}")
# conditional formatting on variance (H) and remaining fee (J)
ws.conditional_formatting.add(f"H{LOG_FIRST}:H{LLAST}",
    CellIsRule(operator="lessThan", formula=["0"], fill=FILL_RED, font=FONT_RED))
ws.conditional_formatting.add(f"H{LOG_FIRST}:H{LLAST}",
    CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_AMBER, font=FONT_AMBER))
ws.conditional_formatting.add(f"J{LOG_FIRST}:J{LLAST}",
    CellIsRule(operator="lessThan", formula=["0"], fill=FILL_RED, font=FONT_RED))
ws.freeze_panes = f"A{LOG_FIRST}"

LOG_F = f"'Forecast Log'!$F${LOG_FIRST}:$F${LLAST}"
LOG_A = f"'Forecast Log'!$A${LOG_FIRST}:$A${LLAST}"
LOG_B = f"'Forecast Log'!$B${LOG_FIRST}:$B${LLAST}"

# ============================================================ DASHBOARD
ws = wb.create_sheet("Dashboard")
ws.sheet_view.showGridLines = False
ws["A1"] = "Portfolio Dashboard"; ws["A1"].font = F_TITLE
ws["A2"] = "Financial year September 2026 – August 2027. Forecast from the Forecast Log; actuals from Month-End Provisions."
ws["A2"].font = F_SUB

# --- monthly table ---
m_hdr_row = 4
hdr(ws, m_hdr_row, ["Month", "Forecast Recognised (£)", "Actual NFI (£)", "Variance (£)"], start=1)
fy_months = [dt.date(2026,9,1)]
for k in range(1,12):
    y = 2026 + (9-1+k)//12
    mth = (9-1+k)%12 + 1
    fy_months.append(dt.date(y, mth, 1))
m_first = m_hdr_row + 1
for i, mdate in enumerate(fy_months):
    rr = m_first + i
    a = ws.cell(row=rr, column=1, value=mdate); a.font = F_BODY; a.number_format = DATEF; a.alignment = CENTER; a.border = BORDER
    b = ws.cell(row=rr, column=2, value=f"=SUMIFS({LOG_F},{LOG_B},$A{rr})")
    b.font = F_BODY; b.number_format = GBP; b.border = BORDER
    c = ws.cell(row=rr, column=3, value=f"=SUMIFS({PROV_NFI},{PROV_MONTH},$A{rr})")
    c.font = F_BODY; c.number_format = GBP; c.border = BORDER
    d = ws.cell(row=rr, column=4, value=f"=$C{rr}-$B{rr}")
    d.font = F_BODY; d.number_format = GBP; d.border = BORDER
m_last = m_first + 11
tr = m_last + 1
tc = ws.cell(row=tr, column=1, value="Total"); tc.font = F_BOLD; tc.fill = FILL_TOTAL; tc.border = BORDER
for col in (2,3,4):
    L = get_column_letter(col)
    cc = ws.cell(row=tr, column=col, value=f"=SUM({L}{m_first}:{L}{m_last})")
    cc.font = F_BOLD; cc.number_format = GBP; cc.fill = FILL_TOTAL; cc.border = BORDER
ws.conditional_formatting.add(f"D{m_first}:D{m_last}",
    CellIsRule(operator="lessThan", formula=["0"], fill=FILL_RED, font=FONT_RED))
ws.conditional_formatting.add(f"D{m_first}:D{m_last}",
    CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_AMBER, font=FONT_AMBER))

# --- by engagement table ---
e_hdr_row = m_hdr_row
ecol = 6  # start column F
eheaders = ["Matter Code", "Client", "Engagement Manager", "Total Fee (£)",
            "Forecast to Date (£)", "Actual to Date (£)", "Remaining Fee (£)", "% Recognised"]
hdr(ws, e_hdr_row, eheaders, start=ecol)
e_first = e_hdr_row + 1
E_ROWS = 20
for i in range(E_ROWS):
    rr = e_first + i
    regrow = REG_FIRST + i
    code = f"$F{rr}"
    a = ws.cell(row=rr, column=ecol, value=f"=IF('Matters Register'!$A{regrow}=\"\",\"\",'Matters Register'!$A{regrow})")
    a.font = F_LINK; a.alignment = CENTER; a.border = BORDER
    cl = ws.cell(row=rr, column=ecol+1, value=f"=IF({code}=\"\",\"\",'Matters Register'!$B{regrow})")
    cl.font = F_LINK; cl.border = BORDER
    em = ws.cell(row=rr, column=ecol+2, value=f"=IF({code}=\"\",\"\",'Matters Register'!$D{regrow})")
    em.font = F_LINK; em.border = BORDER
    fee = ws.cell(row=rr, column=ecol+3, value=f"=IF({code}=\"\",\"\",'Matters Register'!$E{regrow})")
    fee.font = F_LINK; fee.number_format = GBP; fee.border = BORDER
    fc = ws.cell(row=rr, column=ecol+4, value=f"=IF({code}=\"\",\"\",SUMIFS({LOG_F},{LOG_A},{code}))")
    fc.font = F_BODY; fc.number_format = GBP; fc.border = BORDER
    ac = ws.cell(row=rr, column=ecol+5, value=f"=IF({code}=\"\",\"\",SUMIFS({PROV_NFI},{PROV_MATTER},{code}))")
    ac.font = F_BODY; ac.number_format = GBP; ac.border = BORDER
    rem = ws.cell(row=rr, column=ecol+6, value=f"=IF({code}=\"\",\"\",{get_column_letter(ecol+3)}{rr}-{get_column_letter(ecol+5)}{rr})")
    rem.font = F_BODY; rem.number_format = GBP; rem.border = BORDER
    pct = ws.cell(row=rr, column=ecol+7,
        value=f"=IF(OR({code}=\"\",{get_column_letter(ecol+3)}{rr}=0),\"\",{get_column_letter(ecol+5)}{rr}/{get_column_letter(ecol+3)}{rr})")
    pct.font = F_BODY; pct.number_format = PCT; pct.alignment = CENTER; pct.border = BORDER
e_last = e_first + E_ROWS - 1
EM_CODE  = f"$F${e_first}:$F${e_last}"
EM_EMCOL = f"${get_column_letter(ecol+2)}${e_first}:${get_column_letter(ecol+2)}${e_last}"
EM_FC    = f"${get_column_letter(ecol+4)}${e_first}:${get_column_letter(ecol+4)}${e_last}"
EM_AC    = f"${get_column_letter(ecol+5)}${e_first}:${get_column_letter(ecol+5)}${e_last}"

# --- by EM table (list of EMs maintained here) ---
emhdr_row = e_last + 3
ws.cell(row=emhdr_row-1, column=ecol, value="By Engagement Manager (add/remove EM names as needed)").font = F_SUB
hdr(ws, emhdr_row, ["Engagement Manager", "Forecast to Date (£)", "Actual to Date (£)"], start=ecol)
em_names = ["Priya Shah", "James Okoro"]
emf = emhdr_row + 1
for i, name in enumerate(em_names):
    rr = emf + i
    n = ws.cell(row=rr, column=ecol, value=name); n.font = F_INPUT; n.fill = FILL_INPUT; n.border = BORDER
    fc = ws.cell(row=rr, column=ecol+1, value=f"=SUMIF({EM_EMCOL},${get_column_letter(ecol)}{rr},{EM_FC})")
    fc.font = F_BODY; fc.number_format = GBP; fc.border = BORDER
    ac = ws.cell(row=rr, column=ecol+2, value=f"=SUMIF({EM_EMCOL},${get_column_letter(ecol)}{rr},{EM_AC})")
    ac.font = F_BODY; ac.number_format = GBP; ac.border = BORDER

# widths
for col, w in zip("ABCD", [12, 22, 16, 14]):
    ws.column_dimensions[col].width = w
ws.column_dimensions["E"].width = 3
for col, w in zip(["F","G","H","I","J","K","L","M"], [12, 20, 20, 14, 16, 16, 16, 12]):
    ws.column_dimensions[col].width = w
ws.row_dimensions[m_hdr_row].height = 30

# --- chart ---
chart = BarChart(); chart.type = "col"; chart.style = 10
chart.title = "Forecast vs Actual NFI by Month (FY Sep 2026 – Aug 2027)"
chart.y_axis.title = "£"; chart.x_axis.title = "Month"
data = Reference(ws, min_col=2, max_col=3, min_row=m_hdr_row, max_row=m_last)
cats = Reference(ws, min_col=1, min_row=m_first, max_row=m_last)
chart.add_data(data, titles_from_data=True); chart.set_categories(cats)
chart.height = 8; chart.width = 20
ws.add_chart(chart, f"A{tr+3}")

# order sheets: Instructions, Register, Log, Provisions, Dashboard
wb.move_sheet("Month-End Provisions", offset=1)  # keep after Log? ensure order below
order = ["Instructions", "Matters Register", "Forecast Log", "Month-End Provisions", "Dashboard"]
wb._sheets.sort(key=lambda s: order.index(s.title))
wb.active = wb.sheetnames.index("Instructions")

out = "Income_Forecast.xlsx"
wb.save(out)
print("saved", out)
