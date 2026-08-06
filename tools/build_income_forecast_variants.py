#!/usr/bin/env python3
"""Build XLOOKUP variants of the income forecast workbook.

Produces two files (both use XLOOKUP throughout, for Excel):
  - Income_Forecast_XLOOKUP.xlsx        : cumulative % complete model (as the main file)
  - Income_Forecast_ForecastAmount.xlsx : engagement manager enters a direct £ forecast
                                          per month; cumulative and % complete are derived.

Run:  python tools/build_income_forecast_variants.py

Note: openpyxl writes formulas without cached results, so a freshly generated file shows
blank formula cells until opened in Excel, which recalculates on open. (The committed
copies also carry verified cached values injected separately.)
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

ARIAL = "Arial"
F_TITLE = Font(name=ARIAL, size=16, bold=True, color="1F3864")
F_SUB   = Font(name=ARIAL, size=10, italic=True, color="595959")
F_HDR   = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_BODY  = Font(name=ARIAL, size=10, color="000000")
F_INPUT = Font(name=ARIAL, size=10, color="0000FF")
F_BOLD  = Font(name=ARIAL, size=10, bold=True, color="000000")
F_LINK  = Font(name=ARIAL, size=10, color="008000")
FILL_HDR   = PatternFill("solid", fgColor="1F3864")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_CALC  = PatternFill("solid", fgColor="F2F2F2")
FILL_TOTAL = PatternFill("solid", fgColor="D9E1F2")
FILL_RED   = PatternFill("solid", fgColor="F8CBAD")
FILL_AMBER = PatternFill("solid", fgColor="FFE699")
FONT_RED   = Font(name=ARIAL, size=10, color="843C0C")
FONT_AMBER = Font(name=ARIAL, size=10, color="806000")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP   = Alignment(horizontal="left", vertical="top", wrap_text=True)
GBP  = '£#,##0;(£#,##0);-'
GBP0 = '£#,##0'
PCT  = '0.0%'
DATEF = 'MMM-YYYY'
HDR_ROW = 4


def hdr(ws, row, headers, start=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start+i, value=h)
        c.font = F_HDR; c.fill = FILL_HDR; c.alignment = CENTER; c.border = BORDER


def build(mode):
    """mode: 'pct' (cumulative % complete input) or 'amount' (direct £ forecast input)."""
    amount = (mode == "amount")
    wb = Workbook()

    # ---------------- Instructions ----------------
    ws = wb.active; ws.title = "Instructions"; ws.sheet_view.showGridLines = False
    ws["A1"] = "Consultancy Income Forecast"; ws["A1"].font = F_TITLE
    ws["A2"] = ("Direct £ forecast per month" if amount else
                "Percentage-of-completion revenue recognition") + \
               " · GBP (£) · Financial year September–August · XLOOKUP formulas"
    ws["A2"].font = F_SUB
    ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 110
    input_line = ("2.  Forecast Log — the working record, one row per engagement per month. You enter the Matter "
                  "Code, the Month, and " +
                  ("the forecast £ amount to recognise that month" if amount else
                   "the cumulative % complete") +
                  ". Everything else calculates. This is the permanent, append-only history — never overwrite a "
                  "past month's row; add a new row for the new month.")
    recog_line = ("c.  Review the Variance column: red = actual came in UNDER your forecast; amber = actual came in "
                  "OVER your forecast.")
    reall_line = ("d.  Re-allocate next month: " +
                  ("enter next month's forecast £ so the remaining fee (see Remaining Fee) is re-spread across the "
                   "remaining months." if amount else
                   "use Cumulative Actual and Remaining Fee to re-base your % complete so the remaining fee is "
                   "re-spread across the remaining months.") +
                  " The workbook flags the variance; the EM decides how to re-spread it.")
    lines = [
        ("How this workbook works", "h"),
        ("Every engagement is confirmed work with a Matter Code and an agreed fixed fee. Each month the engagement "
         "manager (EM) " + ("enters the £ amount of fee to recognise" if amount else "sets the cumulative % complete") +
         "; the workbook " + ("tracks it against the fee" if amount else "recognises the corresponding slice of the fee") +
         ". Actual net fee income (NFI) is pulled automatically from the Month-End Provisions tab and compared "
         "against the forecast, and any variance is flagged.", "p"),
        ("The five tabs", "h"),
        ("1.  Matters Register — master list of engagements (Matter Code, client, EM, total fixed fee, dates, status). "
         "Add a row here first for every new confirmed engagement.", "p"),
        (input_line, "p"),
        ("3.  Month-End Provisions — paste your monthly NFI provisions extract here (Matter Code, Month, NFI Amount). "
         "The Forecast Log reads actuals from this tab by matching Matter Code + Month.", "p"),
        ("4.  Dashboard — portfolio roll-up by month, by engagement and by EM, with a forecast-vs-actual chart and an "
         "Open Variances tracker.", "p"),
        ("5.  Instructions — this page.", "p"),
        ("Monthly routine", "h"),
        ("a.  Paste the latest NFI provisions into Month-End Provisions.", "p"),
        ("b.  In the Forecast Log, add one new row per active engagement for the new month and enter " +
         ("the forecast £ amount." if amount else "the cumulative % complete."), "p"),
        (recog_line, "p"),
        (reall_line, "p"),
        ("e.  Record the reforecast on the FLAGGED month's row: set Variance Status (Reforecast / Accepted), plus "
         "Reforecast By and Reforecast Date. This is the audit trail.", "p"),
        ("Tracking reforecasts", "h"),
        ("The Forecast Log's right-hand columns show whether each flagged variance has been actioned: Reforecast "
         "Entered? auto-detects whether a later month's row exists; Outstanding? shows YES (red) for any variance not "
         "yet signed off; Action Note says what's needed. Use the header AutoFilter to filter Outstanding? = YES or "
         "by Engagement Manager. The Dashboard's Open Variances panel counts outstanding items overall and per EM.", "p"),
        ("Colour key", "h"),
        ("Blue text on pale-yellow fill = cells you fill in.   Black = calculated, do not edit.   Green = pulled from "
         "another tab.   Red / amber fills = under / over variance and outstanding-reforecast flags.", "p"),
        ("Notes", "h"),
        (("· Forecast and all amounts are GBP.   " if amount else
          "· % complete is stored as a fraction (enter 25% or 0.25, not 25).   · All amounts are GBP.   ") +
         "· Dates are the first of the month (shown as Sep-2026).   "
         "· Demo engagements M-1001/1002/1003 are examples — delete them once you add your own.", "p"),
    ]
    r = 4
    for text, kind in lines:
        c = ws.cell(row=r, column=2, value=text)
        if kind == "h":
            c.font = Font(name=ARIAL, size=12, bold=True, color="1F3864")
        else:
            c.font = F_BODY; c.alignment = WRAP
            ws.row_dimensions[r].height = 15 * (1 + len(text)//95)
        r += 1

    # ---------------- Matters Register ----------------
    ws = wb.create_sheet("Matters Register"); ws.sheet_view.showGridLines = False
    ws["A1"] = "Matters Register"; ws["A1"].font = F_TITLE
    ws["A2"] = "One row per confirmed engagement. Blue cells are inputs."; ws["A2"].font = F_SUB
    reg_headers = ["Matter Code", "Client", "Engagement Name", "Engagement Manager",
                   "Total Fee (£)", "Start Date", "Expected End Date", "Status"]
    hdr(ws, HDR_ROW, reg_headers)
    reg_rows = [
        ["M-1001", "Aurora Retail Group", "Supply Chain Diagnostic", "Priya Shah", 120000,
         dt.date(2026,9,1), dt.date(2027,2,28), "Active"],
        ["M-1002", "Northwind Energy", "Cost Transformation", "James Okoro", 90000,
         dt.date(2026,10,1), dt.date(2027,3,31), "Active"],
        ["M-1003", "Meridian Health", "Operating Model Review", "Priya Shah", 60000,
         dt.date(2026,9,1), dt.date(2026,12,31), "Active"],
    ]
    REG_FIRST = HDR_ROW + 1; REG_CAP = 40
    for i in range(20):
        rr = REG_FIRST + i
        row = reg_rows[i] if i < len(reg_rows) else [None]*8
        for j, val in enumerate(row):
            c = ws.cell(row=rr, column=1+j, value=val)
            c.font = F_INPUT; c.fill = FILL_INPUT; c.border = BORDER
            if j == 4: c.number_format = GBP0
            if j in (5, 6): c.number_format = 'DD-MMM-YYYY'; c.alignment = CENTER
            if j == 0: c.alignment = CENTER
    for i, w in enumerate([12,22,26,22,14,14,16,12]):
        ws.column_dimensions[get_column_letter(1+i)].width = w
    dv = DataValidation(type="list", formula1='"Active,Complete,On hold"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"H{REG_FIRST}:H{REG_FIRST+19}")
    _t = Table(displayName="tblMatters", ref=f"A{HDR_ROW}:H{REG_FIRST+19}")
    _t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=False)
    ws.add_table(_t)
    REG_MATTER = f"'Matters Register'!$A${REG_FIRST}:$A${REG_FIRST+REG_CAP-1}"
    REG_FEE    = f"'Matters Register'!$E${REG_FIRST}:$E${REG_FIRST+REG_CAP-1}"
    REG_EM     = f"'Matters Register'!$D${REG_FIRST}:$D${REG_FIRST+REG_CAP-1}"

    # ---------------- Month-End Provisions ----------------
    ws = wb.create_sheet("Month-End Provisions"); ws.sheet_view.showGridLines = False
    ws["A1"] = "Month-End Provisions (Net Fee Income)"; ws["A1"].font = F_TITLE
    ws["A2"] = "Paste your monthly NFI provisions extract here."; ws["A2"].font = F_SUB
    hdr(ws, HDR_ROW, ["Matter Code", "Month", "NFI Amount (£)"])
    prov_rows = [
        ["M-1001", dt.date(2026,9,1), 12000], ["M-1001", dt.date(2026,10,1), 15000],
        ["M-1001", dt.date(2026,11,1), 26000], ["M-1002", dt.date(2026,10,1), 13500],
        ["M-1003", dt.date(2026,9,1), 12000], ["M-1003", dt.date(2026,10,1), 11000],
    ]
    PROV_FIRST = HDR_ROW + 1; PROV_CAP = 1000
    for i in range(24):
        rr = PROV_FIRST + i
        row = prov_rows[i] if i < len(prov_rows) else [None]*3
        for j, val in enumerate(row):
            c = ws.cell(row=rr, column=1+j, value=val)
            c.font = F_INPUT; c.fill = FILL_INPUT; c.border = BORDER
            if j == 0: c.alignment = CENTER
            if j == 1: c.number_format = DATEF; c.alignment = CENTER
            if j == 2: c.number_format = GBP0
    for col, w in zip("ABC", [12,14,16]):
        ws.column_dimensions[col].width = w
    dvp = DataValidation(type="list", formula1=f"={REG_MATTER}", allow_blank=True)
    ws.add_data_validation(dvp); dvp.add(f"A{PROV_FIRST}:A{PROV_FIRST+23}")
    _t2 = Table(displayName="tblProvisions", ref=f"A{HDR_ROW}:C{PROV_FIRST+23}")
    _t2.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=False)
    ws.add_table(_t2)
    PLAST = PROV_FIRST + PROV_CAP - 1
    PROV_MATTER = f"'Month-End Provisions'!$A${PROV_FIRST}:$A${PLAST}"
    PROV_MONTH  = f"'Month-End Provisions'!$B${PROV_FIRST}:$B${PLAST}"
    PROV_NFI    = f"'Month-End Provisions'!$C${PROV_FIRST}:$C${PLAST}"

    # ---------------- Forecast Log ----------------
    ws = wb.create_sheet("Forecast Log"); ws.sheet_view.showGridLines = False
    ws["A1"] = "Forecast Log — permanent monthly record"; ws["A1"].font = F_TITLE
    ws["A2"] = ("Append one row per engagement per month. Enter Matter Code, Month and " +
                ("Forecast This Month (£)" if amount else "Cumulative % Complete") +
                " (blue). Everything else calculates.")
    ws["A2"].font = F_SUB
    if amount:
        d_hdr, e_hdr, f_hdr = "Forecast This Month (£)", "Forecast Cum to Date (£)", "% Complete to Date"
    else:
        d_hdr, e_hdr, f_hdr = "Cum % Complete", "Forecast Cum Revenue (£)", "Forecast Recognised This Month (£)"
    log_headers = ["Matter Code", "Month", "Total Fee (£)", d_hdr, e_hdr, f_hdr,
                   "Actual NFI This Month (£)", "Variance (£)", "Cumulative Actual to Date (£)",
                   "Remaining Fee (£)", "Engagement Manager", "Variance Status", "Reforecast By",
                   "Reforecast Date", "Reforecast Entered?", "Outstanding?", "Action Note"]
    hdr(ws, HDR_ROW, log_headers)
    LOG_FIRST = HDR_ROW + 1; LOG_N = 60; LLAST = LOG_FIRST + LOG_N - 1
    # demo D values reproduce identical recognised amounts in both models
    if amount:
        d_demo = [12000, 18000, 24000, 13500, 13500, 12000, 12000, 15000]
    else:
        d_demo = [0.10, 0.25, 0.45, 0.15, 0.30, 0.20, 0.40, 0.65]
    log_rows = [
        ["M-1001", dt.date(2026,9,1)], ["M-1001", dt.date(2026,10,1)], ["M-1001", dt.date(2026,11,1)],
        ["M-1002", dt.date(2026,10,1)], ["M-1002", dt.date(2026,11,1)],
        ["M-1003", dt.date(2026,9,1)], ["M-1003", dt.date(2026,10,1)], ["M-1003", dt.date(2026,11,1)],
    ]
    demo_status = {("M-1003", dt.date(2026,10,1)): ("Reforecast", "PS", dt.date(2026,11,4))}
    A_rng = f"$A${LOG_FIRST}:$A${LLAST}"; B_rng = f"$B${LOG_FIRST}:$B${LLAST}"
    D_rng = f"$D${LOG_FIRST}:$D${LLAST}"
    FCOL = "D" if amount else "F"   # column holding "forecast recognised this month"

    for i in range(LOG_N):
        rr = LOG_FIRST + i
        demo = log_rows[i] if i < len(log_rows) else None
        dval = d_demo[i] if i < len(d_demo) else None
        a = ws.cell(row=rr, column=1, value=(demo[0] if demo else None))
        a.font = F_INPUT; a.fill = FILL_INPUT; a.alignment = CENTER; a.border = BORDER
        b = ws.cell(row=rr, column=2, value=(demo[1] if demo else None))
        b.font = F_INPUT; b.fill = FILL_INPUT; b.number_format = DATEF; b.alignment = CENTER; b.border = BORDER
        d = ws.cell(row=rr, column=4, value=dval)
        d.font = F_INPUT; d.fill = FILL_INPUT; d.alignment = CENTER; d.border = BORDER
        d.number_format = GBP0 if amount else PCT
        # C Total Fee via XLOOKUP
        c = ws.cell(row=rr, column=3, value=f'=IF($A{rr}="","",XLOOKUP($A{rr},{REG_MATTER},{REG_FEE},"?"))')
        c.font = F_LINK; c.number_format = GBP0; c.border = BORDER
        if amount:
            # E Forecast cum to date = running sum of forecast £
            e = ws.cell(row=rr, column=5,
                value=f'=IF($A{rr}="","",SUMIFS({D_rng},{A_rng},$A{rr},{B_rng},"<="&$B{rr}))')
            e.font = F_BODY; e.number_format = GBP; e.border = BORDER; e.fill = FILL_CALC
            # F % complete to date = cum / fee
            f = ws.cell(row=rr, column=6,
                value=f'=IF(OR($A{rr}="",$C{rr}=0,$C{rr}="?"),"",$E{rr}/$C{rr})')
            f.font = F_BODY; f.number_format = PCT; f.alignment = CENTER; f.border = BORDER; f.fill = FILL_CALC
        else:
            # E Forecast cum revenue = % x fee
            e = ws.cell(row=rr, column=5,
                value=f'=IF(OR($A{rr}="",$D{rr}="",$C{rr}="?"),"",$D{rr}*$C{rr})')
            e.font = F_BODY; e.number_format = GBP; e.border = BORDER; e.fill = FILL_CALC
            # F recognised this month = cum this month - prior month's cum (reverse XLOOKUP, expanding range)
            if i == 0:
                fval = f'=IF($E{rr}="","",$E{rr})'
            else:
                fval = (f'=IF($E{rr}="","",$E{rr}-IFERROR(XLOOKUP($A{rr},$A${LOG_FIRST}:$A{rr-1},'
                        f'$E${LOG_FIRST}:$E{rr-1},0,0,-1),0))')
            f = ws.cell(row=rr, column=6, value=fval)
            f.font = F_BODY; f.number_format = GBP; f.border = BORDER; f.fill = FILL_CALC
        # G Actual NFI this month (from provisions; blank if none)
        g = ws.cell(row=rr, column=7,
            value=(f'=IF($A{rr}="","",IF(COUNTIFS({PROV_MATTER},$A{rr},{PROV_MONTH},$B{rr})=0,"",'
                   f'SUMIFS({PROV_NFI},{PROV_MATTER},$A{rr},{PROV_MONTH},$B{rr})))'))
        g.font = F_LINK; g.number_format = GBP; g.border = BORDER; g.fill = FILL_CALC
        # H Variance = actual - forecast recognised this month
        h = ws.cell(row=rr, column=8,
            value=f'=IF(OR($G{rr}="",${FCOL}{rr}=""),"",$G{rr}-${FCOL}{rr})')
        h.font = F_BODY; h.number_format = GBP; h.border = BORDER
        # I Cumulative actual to date
        ii = ws.cell(row=rr, column=9,
            value=f'=IF($A{rr}="","",SUMIFS({PROV_NFI},{PROV_MATTER},$A{rr},{PROV_MONTH},"<="&$B{rr}))')
        ii.font = F_LINK; ii.number_format = GBP; ii.border = BORDER; ii.fill = FILL_CALC
        # J Remaining fee = fee - cumulative actual
        jj = ws.cell(row=rr, column=10, value=f'=IF(OR($A{rr}="",$C{rr}="?"),"",$C{rr}-$I{rr})')
        jj.font = F_BODY; jj.number_format = GBP; jj.border = BORDER; jj.fill = FILL_CALC
        # K Engagement Manager via XLOOKUP
        k = ws.cell(row=rr, column=11, value=f'=IF($A{rr}="","",XLOOKUP($A{rr},{REG_MATTER},{REG_EM},""))')
        k.font = F_LINK; k.border = BORDER
        dstat = demo_status.get((demo[0], demo[1])) if demo else None
        l = ws.cell(row=rr, column=12, value=(dstat[0] if dstat else None))
        l.font = F_INPUT; l.fill = FILL_INPUT; l.alignment = CENTER; l.border = BORDER
        m = ws.cell(row=rr, column=13, value=(dstat[1] if dstat else None))
        m.font = F_INPUT; m.fill = FILL_INPUT; m.alignment = CENTER; m.border = BORDER
        n = ws.cell(row=rr, column=14, value=(dstat[2] if dstat else None))
        n.font = F_INPUT; n.fill = FILL_INPUT; n.number_format = 'DD-MMM-YYYY'; n.alignment = CENTER; n.border = BORDER
        o = ws.cell(row=rr, column=15,
            value=(f'=IF(OR($H{rr}="",$H{rr}=0),"",IF(COUNTIFS({A_rng},$A{rr},{B_rng},">"&$B{rr})>0,"Yes","No"))'))
        o.font = F_BODY; o.alignment = CENTER; o.border = BORDER; o.fill = FILL_CALC
        p = ws.cell(row=rr, column=16,
            value=f'=IF(AND($H{rr}<>"",$H{rr}<>0,NOT(OR($L{rr}="Reforecast",$L{rr}="Accepted"))),"YES","")')
        p.font = F_BOLD; p.alignment = CENTER; p.border = BORDER; p.fill = FILL_CALC
        q = ws.cell(row=rr, column=17,
            value=(f'=IF(OR($H{rr}="",$H{rr}=0),"",IF(OR($L{rr}="Reforecast",$L{rr}="Accepted"),"Signed off: "&$L{rr},'
                   f'IF($O{rr}="Yes","Reforecast entered — confirm sign-off","Awaiting reforecast")))'))
        q.font = F_BODY; q.border = BORDER; q.fill = FILL_CALC

    for i, w in enumerate([12,11,13,14,16,18,16,12,16,14,18,13,12,14,15,12,32]):
        ws.column_dimensions[get_column_letter(1+i)].width = w
    ws.row_dimensions[HDR_ROW].height = 42
    for cc in range(1, len(log_headers)+1):
        ws.cell(row=HDR_ROW, column=cc).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    dv_lm = DataValidation(type="list", formula1=f"={REG_MATTER}", allow_blank=True)
    ws.add_data_validation(dv_lm); dv_lm.add(f"A{LOG_FIRST}:A{LLAST}")
    if amount:
        dv_d = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                              error="Enter the forecast fee to recognise this month (£).", errorTitle="Forecast £")
    else:
        dv_d = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                              error="Enter cumulative % complete as a fraction 0–1 (e.g. 0.25).", errorTitle="% complete")
    ws.add_data_validation(dv_d); dv_d.add(f"D{LOG_FIRST}:D{LLAST}")
    dv_st = DataValidation(type="list", formula1='"Open,Reforecast,Accepted"', allow_blank=True)
    ws.add_data_validation(dv_st); dv_st.add(f"L{LOG_FIRST}:L{LLAST}")
    ws.conditional_formatting.add(f"H{LOG_FIRST}:H{LLAST}",
        CellIsRule(operator="lessThan", formula=["0"], fill=FILL_RED, font=FONT_RED))
    ws.conditional_formatting.add(f"H{LOG_FIRST}:H{LLAST}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_AMBER, font=FONT_AMBER))
    ws.conditional_formatting.add(f"J{LOG_FIRST}:J{LLAST}",
        CellIsRule(operator="lessThan", formula=["0"], fill=FILL_RED, font=FONT_RED))
    ws.conditional_formatting.add(f"P{LOG_FIRST}:P{LLAST}",
        CellIsRule(operator="equal", formula=['"YES"'], fill=FILL_RED, font=FONT_RED))
    ws.conditional_formatting.add(f"Q{LOG_FIRST}:Q{LLAST}",
        FormulaRule(formula=[f'AND($Q{LOG_FIRST}<>"",LEFT($Q{LOG_FIRST},6)<>"Signed")'], fill=FILL_AMBER, font=FONT_AMBER))
    ws.auto_filter.ref = f"A{HDR_ROW}:Q{LLAST}"
    ws.freeze_panes = f"C{LOG_FIRST}"

    FCOL_RNG = f"'Forecast Log'!${FCOL}${LOG_FIRST}:${FCOL}${LLAST}"
    LOG_A = f"'Forecast Log'!$A${LOG_FIRST}:$A${LLAST}"
    LOG_B = f"'Forecast Log'!$B${LOG_FIRST}:$B${LLAST}"
    LOG_H = f"'Forecast Log'!$H${LOG_FIRST}:$H${LLAST}"
    LOG_K = f"'Forecast Log'!$K${LOG_FIRST}:$K${LLAST}"
    LOG_P = f"'Forecast Log'!$P${LOG_FIRST}:$P${LLAST}"

    # ---------------- Dashboard ----------------
    ws = wb.create_sheet("Dashboard"); ws.sheet_view.showGridLines = False
    ws["A1"] = "Portfolio Dashboard"; ws["A1"].font = F_TITLE
    ws["A2"] = "Financial year September 2026 – August 2027."; ws["A2"].font = F_SUB
    m_hdr = HDR_ROW
    hdr(ws, m_hdr, ["Month", "Forecast (£)", "Actual NFI (£)", "Variance (£)"])
    fy = [dt.date(2026,9,1)]
    for k in range(1,12):
        y = 2026 + (8+k)//12; mth = (8+k)%12 + 1; fy.append(dt.date(y, mth, 1))
    m_first = m_hdr + 1
    for i, md in enumerate(fy):
        rr = m_first + i
        a = ws.cell(row=rr, column=1, value=md); a.font = F_BODY; a.number_format = DATEF; a.alignment = CENTER; a.border = BORDER
        b = ws.cell(row=rr, column=2, value=f"=SUMIFS({FCOL_RNG},{LOG_B},$A{rr})"); b.font = F_BODY; b.number_format = GBP; b.border = BORDER
        c = ws.cell(row=rr, column=3, value=f"=SUMIFS({PROV_NFI},{PROV_MONTH},$A{rr})"); c.font = F_BODY; c.number_format = GBP; c.border = BORDER
        d = ws.cell(row=rr, column=4, value=f"=$C{rr}-$B{rr}"); d.font = F_BODY; d.number_format = GBP; d.border = BORDER
    m_last = m_first + 11; tr = m_last + 1
    tc = ws.cell(row=tr, column=1, value="Total"); tc.font = F_BOLD; tc.fill = FILL_TOTAL; tc.border = BORDER
    for col in (2,3,4):
        L = get_column_letter(col)
        cc = ws.cell(row=tr, column=col, value=f"=SUM({L}{m_first}:{L}{m_last})")
        cc.font = F_BOLD; cc.number_format = GBP; cc.fill = FILL_TOTAL; cc.border = BORDER
    ws.conditional_formatting.add(f"D{m_first}:D{m_last}",
        CellIsRule(operator="lessThan", formula=["0"], fill=FILL_RED, font=FONT_RED))
    ws.conditional_formatting.add(f"D{m_first}:D{m_last}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_AMBER, font=FONT_AMBER))

    ecol = 6; e_hdr_row = m_hdr
    eheaders = ["Matter Code", "Client", "Engagement Manager", "Total Fee (£)", "Forecast to Date (£)",
                "Actual to Date (£)", "Remaining Fee (£)", "% Recognised", "Open Var (#)"]
    hdr(ws, e_hdr_row, eheaders, start=ecol)
    e_first = e_hdr_row + 1; E_ROWS = 20
    for i in range(E_ROWS):
        rr = e_first + i; regrow = REG_FIRST + i; code = f"$F{rr}"
        gc = lambda off: get_column_letter(ecol+off)
        ws.cell(row=rr, column=ecol, value=f"=IF('Matters Register'!$A{regrow}=\"\",\"\",'Matters Register'!$A{regrow})").font = F_LINK
        ws.cell(row=rr, column=ecol).alignment = CENTER; ws.cell(row=rr, column=ecol).border = BORDER
        for off, formula, fmt in [
            (1, f"=IF({code}=\"\",\"\",XLOOKUP({code},{REG_MATTER},'Matters Register'!$B${REG_FIRST}:$B${REG_FIRST+REG_CAP-1},\"\"))", None),
            (2, f"=IF({code}=\"\",\"\",XLOOKUP({code},{REG_MATTER},{REG_EM},\"\"))", None),
            (3, f"=IF({code}=\"\",\"\",XLOOKUP({code},{REG_MATTER},{REG_FEE},\"\"))", GBP),
            (4, f"=IF({code}=\"\",\"\",SUMIFS({FCOL_RNG},{LOG_A},{code}))", GBP),
            (5, f"=IF({code}=\"\",\"\",SUMIFS({PROV_NFI},{PROV_MATTER},{code}))", GBP),
            (6, f"=IF({code}=\"\",\"\",{gc(3)}{rr}-{gc(5)}{rr})", GBP),
            (7, f"=IF(OR({code}=\"\",{gc(3)}{rr}=0),\"\",{gc(5)}{rr}/{gc(3)}{rr})", PCT),
            (8, f"=IF({code}=\"\",\"\",COUNTIFS({LOG_A},{code},{LOG_P},\"YES\"))", None),
        ]:
            cell = ws.cell(row=rr, column=ecol+off, value=formula)
            cell.font = F_BODY if off >= 4 else F_LINK; cell.border = BORDER
            if fmt: cell.number_format = fmt
            if off in (7, 8): cell.alignment = CENTER
    e_last = e_first + E_ROWS - 1
    ws.conditional_formatting.add(f"{get_column_letter(ecol+8)}{e_first}:{get_column_letter(ecol+8)}{e_last}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_RED, font=FONT_RED))
    EM_EMCOL = f"${get_column_letter(ecol+2)}${e_first}:${get_column_letter(ecol+2)}${e_last}"
    EM_FC    = f"${get_column_letter(ecol+4)}${e_first}:${get_column_letter(ecol+4)}${e_last}"
    EM_AC    = f"${get_column_letter(ecol+5)}${e_first}:${get_column_letter(ecol+5)}${e_last}"

    emhdr = e_last + 3
    ws.cell(row=emhdr-1, column=ecol, value="By Engagement Manager (add/remove names as needed)").font = F_SUB
    hdr(ws, emhdr, ["Engagement Manager", "Forecast to Date (£)", "Actual to Date (£)", "Open (#)"], start=ecol)
    em_names = ["Priya Shah", "James Okoro"]; emf = emhdr + 1
    for i, name in enumerate(em_names):
        rr = emf + i
        n = ws.cell(row=rr, column=ecol, value=name); n.font = F_INPUT; n.fill = FILL_INPUT; n.border = BORDER
        self_c = f"${get_column_letter(ecol)}{rr}"
        fc = ws.cell(row=rr, column=ecol+1, value=f"=SUMIF({EM_EMCOL},{self_c},{EM_FC})"); fc.font = F_BODY; fc.number_format = GBP; fc.border = BORDER
        ac = ws.cell(row=rr, column=ecol+2, value=f"=SUMIF({EM_EMCOL},{self_c},{EM_AC})"); ac.font = F_BODY; ac.number_format = GBP; ac.border = BORDER
        ov = ws.cell(row=rr, column=ecol+3, value=f"=COUNTIFS({LOG_K},{self_c},{LOG_P},\"YES\")"); ov.font = F_BODY; ov.alignment = CENTER; ov.border = BORDER
    emlast = emf + len(em_names) - 1
    ws.conditional_formatting.add(f"{get_column_letter(ecol+3)}{emf}:{get_column_letter(ecol+3)}{emlast}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_RED, font=FONT_RED))

    for col, w in zip("ABCD", [20,22,16,14]): ws.column_dimensions[col].width = w
    ws.column_dimensions["E"].width = 3
    for col, w in zip(["F","G","H","I","J","K","L","M","N"], [12,20,18,11,16,16,16,12,12]):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[m_hdr].height = 30

    ov_row = tr + 2
    ws.cell(row=ov_row, column=1, value="Open Variances — flagged but not yet reforecast").font = \
        Font(name=ARIAL, size=11, bold=True, color="1F3864")
    panel = [("Open items (#)", f'=COUNTIF({LOG_P},"YES")', None),
             ("Open value (£, net)", f'=SUMIFS({LOG_H},{LOG_P},"YES")', GBP),
             ("Oldest open month", f'=IF(COUNTIF({LOG_P},"YES")=0,"",_xlfn.MINIFS({LOG_B},{LOG_P},"YES"))', DATEF)]
    for j, (label, formula, fmt) in enumerate(panel):
        r2 = ov_row + 1 + j
        lc = ws.cell(row=r2, column=1, value=label); lc.font = F_BODY; lc.border = BORDER
        vc = ws.cell(row=r2, column=2, value=formula); vc.font = F_BOLD; vc.border = BORDER; vc.alignment = CENTER
        if fmt: vc.number_format = fmt
    ws.conditional_formatting.add(f"B{ov_row+1}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_RED, font=FONT_RED))

    chart = BarChart(); chart.type = "col"; chart.style = 10
    chart.title = "Forecast vs Actual NFI by Month (FY Sep 2026 – Aug 2027)"
    chart.y_axis.title = "£"; chart.x_axis.title = "Month"
    data = Reference(ws, min_col=2, max_col=3, min_row=m_hdr, max_row=m_last)
    cats = Reference(ws, min_col=1, min_row=m_first, max_row=m_last)
    chart.add_data(data, titles_from_data=True); chart.set_categories(cats)
    chart.height = 8; chart.width = 20
    ws.add_chart(chart, f"A{ov_row+5}")

    order = ["Instructions", "Matters Register", "Forecast Log", "Month-End Provisions", "Dashboard"]
    wb._sheets.sort(key=lambda s: order.index(s.title))
    wb.active = wb.sheetnames.index("Instructions")
    return wb


if __name__ == "__main__":
    for mode, out in [("pct", "Income_Forecast_XLOOKUP.xlsx"),
                      ("amount", "Income_Forecast_ForecastAmount.xlsx")]:
        build(mode).save(out); print("saved", out)
